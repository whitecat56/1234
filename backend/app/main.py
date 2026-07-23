import os, shutil, tempfile
from datetime import datetime, date, timedelta, timezone
from pathlib import Path
from typing import Optional, List
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordBearer
from jose import jwt, JWTError
from passlib.context import CryptContext
from pydantic import BaseModel, Field
from sqlalchemy import create_engine, Column, Integer, String, Float, DateTime, ForeignKey, Text, func, Date, or_
from sqlalchemy.orm import declarative_base, sessionmaker, relationship, Session
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
DB_PATH = BASE_DIR / 'gameclub_aktash.db'
SECRET_KEY = os.getenv('SECRET_KEY', 'change-this-local-secret-key')
ALGORITHM = 'HS256'
ACCESS_TOKEN_EXPIRE_MINUTES = 720
engine = create_engine(f'sqlite:///{DB_PATH}', connect_args={'check_same_thread': False})
SessionLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
Base = declarative_base()
pwd_context = CryptContext(schemes=['bcrypt'], deprecated='auto')
oauth2_scheme = OAuth2PasswordBearer(tokenUrl='/api/login')

class User(Base):
    __tablename__='users'; id=Column(Integer, primary_key=True); username=Column(String, unique=True, index=True); full_name=Column(String); role=Column(String, default='employee'); hashed_password=Column(String); is_active=Column(Integer, default=1); created_at=Column(DateTime, default=datetime.utcnow)
class Product(Base):
    __tablename__='products'; id=Column(Integer, primary_key=True); image=Column(Text, default=''); barcode=Column(String, nullable=True, index=True); name=Column(String, index=True); category=Column(String, index=True); purchase_price=Column(Float); selling_price=Column(Float); quantity=Column(Integer); description=Column(Text, default=''); created_at=Column(DateTime, default=datetime.utcnow); updated_at=Column(DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
class Sale(Base):
    __tablename__='sales'; id=Column(Integer, primary_key=True); user_id=Column(Integer, ForeignKey('users.id')); total_revenue=Column(Float); total_profit=Column(Float); created_at=Column(DateTime, default=datetime.utcnow); user=relationship('User'); items=relationship('SaleItem', cascade='all, delete-orphan')
class SaleItem(Base):
    __tablename__='sale_items'; id=Column(Integer, primary_key=True); sale_id=Column(Integer, ForeignKey('sales.id')); product_id=Column(Integer, ForeignKey('products.id')); product_name=Column(String); quantity=Column(Integer); unit_purchase_price=Column(Float); unit_selling_price=Column(Float); revenue=Column(Float); profit=Column(Float); product=relationship('Product')
class History(Base):
    __tablename__='history'; id=Column(Integer, primary_key=True); user_id=Column(Integer, nullable=True); username=Column(String); action=Column(String, index=True); entity=Column(String); entity_id=Column(Integer, nullable=True); details=Column(Text, default=''); created_at=Column(DateTime, default=datetime.utcnow)
class Setting(Base):
    __tablename__='settings'; key=Column(String, primary_key=True); value=Column(Text)
class Expense(Base):
    __tablename__='expenses'; id=Column(Integer, primary_key=True); title=Column(String); amount=Column(Float); created_at=Column(DateTime, default=datetime.utcnow)

class LoginIn(BaseModel): username:str; password:str
class Token(BaseModel): access_token:str; token_type:str='bearer'; user:dict
class UserIn(BaseModel): username:str; full_name:str; password:str=Field(min_length=4); role:str='employee'; is_active:int=1
class UserOut(BaseModel): id:int; username:str; full_name:str; role:str; is_active:int; created_at:datetime
class UserUpdate(BaseModel): full_name:Optional[str]=None; role:Optional[str]=None; is_active:Optional[int]=None; password:Optional[str]=None
class ProductIn(BaseModel): image:str=''; barcode:Optional[str]=None; name:str; category:str; purchase_price:float=Field(ge=0); selling_price:float=Field(ge=0); quantity:int=Field(ge=0); description:str=''
class ProductOut(ProductIn): id:int; created_at:datetime; updated_at:datetime
class SaleLine(BaseModel): product_id:int; quantity:int=Field(gt=0)
class SaleIn(BaseModel): items:List[SaleLine]
class ExpenseIn(BaseModel): title:str; amount:float=Field(ge=0)
class SettingsIn(BaseModel): gameclub_name:str='GameClub Aktash'; currency:str='KGS'

app=FastAPI(title='GameClub Aktash Manager API')
app.add_middleware(CORSMiddleware, allow_origins=['http://localhost:5173','http://127.0.0.1:5173'], allow_credentials=True, allow_methods=['*'], allow_headers=['*'])

def db():
    s=SessionLocal();
    try: yield s
    finally: s.close()
def hash_pw(p): return pwd_context.hash(p)
def verify(p,h): return pwd_context.verify(p,h)
def log(s:Session, user, action, entity, entity_id=None, details=''):
    s.add(History(user_id=getattr(user,'id',None), username=getattr(user,'username','system'), action=action, entity=entity, entity_id=entity_id, details=details))
def create_token(data):
    d=data.copy(); d['exp']=datetime.now(timezone.utc)+timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES); return jwt.encode(d, SECRET_KEY, algorithm=ALGORITHM)
def current_user(token:str=Depends(oauth2_scheme), s:Session=Depends(db)):
    try: payload=jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM]); username=payload.get('sub')
    except JWTError: raise HTTPException(401,'Invalid token')
    u=s.query(User).filter(User.username==username, User.is_active==1).first()
    if not u: raise HTTPException(401,'Inactive or missing user')
    return u
def admin(u:User=Depends(current_user)):
    if u.role!='administrator': raise HTTPException(403,'Administrator only')
    return u

def init():
    Base.metadata.create_all(bind=engine); s=SessionLocal()
    if not s.query(User).filter_by(username='admin').first(): s.add(User(username='admin', full_name='Administrator', role='administrator', hashed_password=hash_pw('admin123')))
    defaults={'gameclub_name':'GameClub Aktash','currency':'KGS'}
    for k,v in defaults.items():
        if not s.get(Setting,k): s.add(Setting(key=k,value=v))
    s.commit(); s.close()
init()

def period_bounds(period='today', custom:Optional[date]=None):
    today=date.today(); target=custom or today
    if period=='yesterday': start=target-timedelta(days=1); end=target
    elif period=='week': start=today-timedelta(days=today.weekday()); end=start+timedelta(days=7)
    elif period=='month': start=today.replace(day=1); end=(start.replace(year=start.year+1, month=1) if start.month==12 else start.replace(month=start.month+1))
    elif period=='year': start=today.replace(month=1,day=1); end=start.replace(year=start.year+1)
    else: start=target; end=target+timedelta(days=1)
    return datetime.combine(start, datetime.min.time()), datetime.combine(end, datetime.min.time())
def metrics(s,start,end):
    sales=s.query(Sale).filter(Sale.created_at>=start, Sale.created_at<end).all(); expenses=s.query(Expense).filter(Expense.created_at>=start, Expense.created_at<end).all()
    return {'revenue':sum(x.total_revenue for x in sales),'profit':sum(x.total_profit for x in sales),'expenses':sum(e.amount for e in expenses),'products_sold':sum(i.quantity for x in sales for i in x.items)}
def serialize_sale(x): return {'id':x.id,'user':x.user.username if x.user else 'unknown','total_revenue':x.total_revenue,'total_profit':x.total_profit,'created_at':x.created_at,'items':[{'product_id':i.product_id,'product_name':i.product_name,'quantity':i.quantity,'revenue':i.revenue,'profit':i.profit} for i in x.items]}

@app.post('/api/login', response_model=Token)
def login(data:LoginIn, s:Session=Depends(db)):
    u=s.query(User).filter_by(username=data.username).first()
    if not u or not verify(data.password,u.hashed_password) or not u.is_active: raise HTTPException(401,'Wrong credentials')
    log(s,u,'Login','users',u.id,'User logged in'); s.commit()
    return {'access_token':create_token({'sub':u.username,'role':u.role}), 'user':{'id':u.id,'username':u.username,'full_name':u.full_name,'role':u.role}}
@app.post('/api/logout')
def logout(u:User=Depends(current_user), s:Session=Depends(db)): log(s,u,'Logout','users',u.id,'User logged out'); s.commit(); return {'ok':True}
@app.get('/api/me')
def me(u:User=Depends(current_user)): return {'id':u.id,'username':u.username,'full_name':u.full_name,'role':u.role}
@app.get('/api/products')
def products(search:str='', category:str='', s:Session=Depends(db), u:User=Depends(current_user)):
    q=s.query(Product)
    if search: q=q.filter(or_(Product.name.ilike(f'%{search}%'), Product.barcode.ilike(f'%{search}%'), Product.category.ilike(f'%{search}%')))
    if category: q=q.filter(Product.category==category)
    return q.order_by(Product.name).all()
@app.post('/api/products')
def add_product(p:ProductIn, s:Session=Depends(db), u:User=Depends(admin)):
    obj=Product(**p.model_dump()); s.add(obj); s.flush(); log(s,u,'Added Product','products',obj.id,obj.name); s.commit(); s.refresh(obj); return obj
@app.put('/api/products/{id}')
def edit_product(id:int,p:ProductIn,s:Session=Depends(db),u:User=Depends(admin)):
    obj=s.get(Product,id); 
    if not obj: raise HTTPException(404,'Product not found')
    for k,v in p.model_dump().items(): setattr(obj,k,v)
    obj.updated_at=datetime.utcnow(); log(s,u,'Edited Product','products',id,obj.name); s.commit(); return obj
@app.delete('/api/products/{id}')
def del_product(id:int,s:Session=Depends(db),u:User=Depends(admin)):
    obj=s.get(Product,id); 
    if not obj: raise HTTPException(404,'Product not found')
    name=obj.name; s.delete(obj); log(s,u,'Deleted Product','products',id,name); s.commit(); return {'ok':True}
@app.post('/api/sales')
def sell(sale:SaleIn,s:Session=Depends(db),u:User=Depends(current_user)):
    tx=Sale(user_id=u.id,total_revenue=0,total_profit=0); s.add(tx); s.flush(); rev=prof=0
    for line in sale.items:
        p=s.get(Product,line.product_id)
        if not p: raise HTTPException(404,f'Product {line.product_id} not found')
        if p.quantity<line.quantity: raise HTTPException(400,f'Not enough stock for {p.name}')
        p.quantity-=line.quantity; r=p.selling_price*line.quantity; pr=(p.selling_price-p.purchase_price)*line.quantity; rev+=r; prof+=pr
        s.add(SaleItem(sale_id=tx.id,product_id=p.id,product_name=p.name,quantity=line.quantity,unit_purchase_price=p.purchase_price,unit_selling_price=p.selling_price,revenue=r,profit=pr))
    tx.total_revenue=rev; tx.total_profit=prof; log(s,u,'Sale','sales',tx.id,f'Sold {sum(i.quantity for i in tx.items)} items for {rev}'); s.commit(); return serialize_sale(tx)
@app.get('/api/sales')
def sales(s:Session=Depends(db), u:User=Depends(current_user)): return [serialize_sale(x) for x in s.query(Sale).order_by(Sale.created_at.desc()).limit(300)]
@app.get('/api/dashboard')
def dashboard(s:Session=Depends(db), u:User=Depends(current_user)):
    st,en=period_bounds('today'); mt=metrics(s,st,en); ms,me=period_bounds('month'); mm=metrics(s,ms,me)
    top=s.query(SaleItem.product_name, func.sum(SaleItem.quantity).label('qty')).join(Sale).filter(Sale.created_at>=st,Sale.created_at<en).group_by(SaleItem.product_name).order_by(func.sum(SaleItem.quantity).desc()).first()
    return {**mt,'remaining_products':s.query(func.sum(Product.quantity)).scalar() or 0,'low_stock':s.query(Product).filter(Product.quantity<=5).all(),'most_popular_product':top.product_name if top else 'None','monthly_revenue':mm['revenue'],'recent_sales':[serialize_sale(x) for x in s.query(Sale).order_by(Sale.created_at.desc()).limit(8)],'daily_chart':chart_series(s,14)}
def chart_series(s,days):
    out=[]
    for i in range(days-1,-1,-1):
        d=date.today()-timedelta(days=i); st,en=period_bounds('today',d); m=metrics(s,st,en); out.append({'date':d.isoformat(),**m})
    return out
@app.get('/api/analytics')
def analytics(s:Session=Depends(db), u:User=Depends(current_user)):
    top=s.query(SaleItem.product_name, func.sum(SaleItem.quantity).label('qty')).group_by(SaleItem.product_name).order_by(func.sum(SaleItem.quantity).desc()).limit(10).all()
    return {'daily':chart_series(s,30),'monthly':monthly_series(s),'top_products':[{'name':x.product_name,'quantity':x.qty} for x in top]}
def monthly_series(s):
    out=[]; today=date.today()
    for i in range(11,-1,-1):
        y=today.year; m=today.month-i
        while m<=0: m+=12; y-=1
        st=date(y,m,1); en=date(y+1,1,1) if m==12 else date(y,m+1,1); out.append({'month':st.strftime('%Y-%m'),**metrics(s,datetime.combine(st,datetime.min.time()),datetime.combine(en,datetime.min.time()))})
    return out
@app.get('/api/calendar')
def calendar(period:str='today', selected:Optional[date]=None, s:Session=Depends(db), u:User=Depends(current_user)):
    st,en=period_bounds(period,selected); sal=s.query(Sale).filter(Sale.created_at>=st,Sale.created_at<en).order_by(Sale.created_at.desc()).all(); top={}
    for x in sal:
        for i in x.items: top[i.product_name]=top.get(i.product_name,0)+i.quantity
    return {**metrics(s,st,en),'transactions':[serialize_sale(x) for x in sal],'best_selling_products':[{'name':k,'quantity':v} for k,v in sorted(top.items(), key=lambda kv:kv[1], reverse=True)[:10]]}
@app.get('/api/history')
def history(search:str='', s:Session=Depends(db), u:User=Depends(current_user)):
    q=s.query(History)
    if search: q=q.filter(or_(History.action.ilike(f'%{search}%'),History.username.ilike(f'%{search}%'),History.details.ilike(f'%{search}%')))
    return q.order_by(History.created_at.desc()).limit(500).all()
@app.get('/api/employees')
def employees(s:Session=Depends(db),u:User=Depends(admin)): return s.query(User).order_by(User.created_at.desc()).all()
@app.post('/api/employees')
def add_employee(data:UserIn,s:Session=Depends(db),u:User=Depends(admin)):
    if s.query(User).filter_by(username=data.username).first(): raise HTTPException(400,'Username exists')
    obj=User(username=data.username,full_name=data.full_name,role=data.role,hashed_password=hash_pw(data.password),is_active=data.is_active); s.add(obj); s.flush(); log(s,u,'Employee Action','users',obj.id,'Created employee'); s.commit(); return obj
@app.put('/api/employees/{id}')
def upd_employee(id:int,data:UserUpdate,s:Session=Depends(db),u:User=Depends(admin)):
    obj=s.get(User,id); 
    if not obj: raise HTTPException(404,'Employee not found')
    for k,v in data.model_dump(exclude_unset=True).items(): setattr(obj,'hashed_password' if k=='password' else k, hash_pw(v) if k=='password' else v)
    log(s,u,'Employee Action','users',id,'Updated employee'); s.commit(); return obj
@app.delete('/api/employees/{id}')
def rm_employee(id:int,s:Session=Depends(db),u:User=Depends(admin)):
    obj=s.get(User,id); 
    if not obj or obj.role=='administrator': raise HTTPException(400,'Cannot delete')
    s.delete(obj); log(s,u,'Employee Action','users',id,'Deleted employee'); s.commit(); return {'ok':True}
@app.get('/api/settings')
def get_settings(s:Session=Depends(db),u:User=Depends(current_user)): return {x.key:x.value for x in s.query(Setting)}
@app.put('/api/settings')
def save_settings(data:SettingsIn,s:Session=Depends(db),u:User=Depends(admin)):
    for k,v in data.model_dump().items(): obj=s.get(Setting,k) or Setting(key=k); obj.value=v; s.merge(obj)
    log(s,u,'Edited Settings','settings',None,'Updated settings'); s.commit(); return data
@app.post('/api/expenses')
def add_expense(data:ExpenseIn,s:Session=Depends(db),u:User=Depends(admin)):
    e=Expense(**data.model_dump()); s.add(e); log(s,u,'Added Expense','expenses',None,data.title); s.commit(); return e
@app.get('/api/search')
def search(q:str,s:Session=Depends(db),u:User=Depends(current_user)):
    return {'products':s.query(Product).filter(Product.name.ilike(f'%{q}%')).limit(10).all(),'history':s.query(History).filter(History.details.ilike(f'%{q}%')).limit(10).all(),'sales':[serialize_sale(x) for x in s.query(Sale).join(SaleItem).filter(SaleItem.product_name.ilike(f'%{q}%')).limit(10)]}
@app.get('/api/settings/backup')
def backup(u:User=Depends(admin)): return Response(DB_PATH.read_bytes(), media_type='application/octet-stream', headers={'Content-Disposition':'attachment; filename=gameclub_aktash_backup.db'})
@app.post('/api/settings/restore')
def restore(file:UploadFile=File(...),u:User=Depends(admin)):
    with tempfile.NamedTemporaryFile(delete=False) as tmp: shutil.copyfileobj(file.file,tmp); tmp_path=tmp.name
    shutil.copy(tmp_path, DB_PATH); return {'ok':True}
@app.get('/api/settings/export')
def export_excel(s:Session=Depends(db),u:User=Depends(admin)):
    wb=Workbook(); ws=wb.active; ws.title='products'; ws.append(['name','category','barcode','purchase_price','selling_price','quantity','description','image'])
    for p in s.query(Product): ws.append([p.name,p.category,p.barcode,p.purchase_price,p.selling_price,p.quantity,p.description,p.image])
    out=tempfile.NamedTemporaryFile(delete=False,suffix='.xlsx'); wb.save(out.name); data=Path(out.name).read_bytes(); return Response(data, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition':'attachment; filename=gameclub_export.xlsx'})
@app.post('/api/settings/import')
def import_excel(file:UploadFile=File(...),s:Session=Depends(db),u:User=Depends(admin)):
    tmp=tempfile.NamedTemporaryFile(delete=False,suffix='.xlsx'); shutil.copyfileobj(file.file,tmp); wb=load_workbook(tmp.name); ws=wb.active; count=0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]: continue
        s.add(Product(name=row[0],category=row[1] or 'General',barcode=row[2],purchase_price=float(row[3] or 0),selling_price=float(row[4] or 0),quantity=int(row[5] or 0),description=row[6] or '',image=row[7] or '')); count+=1
    log(s,u,'Import Excel','products',None,f'Imported {count} products'); s.commit(); return {'imported':count}
